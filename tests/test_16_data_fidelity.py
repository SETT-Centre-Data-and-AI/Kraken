from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock

import pandas as pd
import pytest
from openpyxl import load_workbook

from kraken import export_results, extract_spreadsheets
from kraken.analysis.data_manipulation import date_converter_df, examine
from kraken.connection.connector import CursorLike, SaConnector
from kraken.support.progress import Progress


def fetch_fidelity_rows(
    rows: list[tuple], columns: list[str], clean_df: bool = True, batch_size: int = 0
) -> pd.DataFrame:
    connector = object.__new__(SaConnector)
    connector.cursor = Mock(spec=CursorLike)
    connector.cursor.description = [
        (column, None, None, None, None, None, None) for column in columns
    ]
    connector.cursor.fetchall.side_effect = [rows, []]
    connector.cursor.fetchmany.side_effect = [rows, []]
    connector.batch_size = batch_size
    connector.multiset_supported = False
    return connector._fetch_data(
        query_name="fidelity",
        progress=Mock(spec=Progress),
        clean_df=clean_df,
        check_column_duplicates=False,
    )[0]


@pytest.mark.parametrize(
    "value, expected_dtype",
    [
        (2**53 + 1, "Int64"),
        (2**63 - 1, "Int64"),
        (2**63, "UInt64"),
        (-(2**63), "Int64"),
        (2**64, "object"),
        (-(2**63) - 1, "object"),
        (1.0, "float64"),
        (True, "boolean"),
        (Decimal("1234567890.123456789"), "object"),
        (date(2026, 1, 1), "object"),
        (datetime(2026, 1, 1), "datetime64[ns]"),
    ],
    ids=[
        "large-integer",
        "signed-limit",
        "unsigned",
        "signed-minimum",
        "huge-positive",
        "huge-negative",
        "float",
        "boolean",
        "decimal",
        "date",
        "datetime",
    ],
)
def test_sql_fidelity_preserves_nullable_scalar_types(
    value: object, expected_dtype: str
) -> None:
    restored = fetch_fidelity_rows([(value,), (None,)], ["value"])

    assert str(restored["value"].dtype) == expected_dtype
    if expected_dtype in ("Int64", "UInt64"):
        assert int(restored["value"].iloc[0]) == value
    else:
        assert restored["value"].iloc[0] == value
    assert pd.isna(restored["value"].iloc[1])


def test_sql_fidelity_preserves_duplicate_columns_and_batches() -> None:
    restored = fetch_fidelity_rows(
        [(2**53 + 1, 1.0), (None, None)], ["value", "value"], batch_size=2
    )

    assert restored.columns.tolist() == ["value", "value"]
    assert int(restored.iloc[0, 0]) == 2**53 + 1
    assert str(restored.iloc[:, 0].dtype) == "Int64"
    assert str(restored.iloc[:, 1].dtype) == "float64"


def test_sql_fidelity_preserves_cleaning_opt_out() -> None:
    restored = fetch_fidelity_rows([(1,), (None,)], ["value"], clean_df=False)

    assert str(restored["value"].dtype) == "float64"


def test_sql_fidelity_keeps_mixed_numeric_scalars_exact() -> None:
    restored = fetch_fidelity_rows([(2**53 + 1,), (1.5,), (None,)], ["value"])

    assert restored["value"].dtype == object
    assert restored["value"].tolist() == [2**53 + 1, 1.5, None]
    assert isinstance(restored["value"].iloc[0], int)


@pytest.mark.parametrize(
    "value, dtype", [(1, "int64"), (1.0, "float64"), (True, "bool")]
)
def test_sql_fidelity_keeps_non_nullable_dtypes(value: object, dtype: str) -> None:
    restored = fetch_fidelity_rows([(value,), (value,)], ["value"])

    assert str(restored["value"].dtype) == dtype


@pytest.mark.parametrize("rows", [[], [(None,), (None,)]], ids=["empty", "all-null"])
def test_sql_fidelity_keeps_untyped_results(rows: list[tuple]) -> None:
    restored = fetch_fidelity_rows(rows, ["value"])

    assert restored.columns.tolist() == ["value"]
    assert restored.shape == (len(rows), 1)
    assert restored["value"].isna().all()


@pytest.mark.parametrize(
    "values",
    [
        [pd.Timestamp("2026-01-01"), pd.NaT],
        [pd.Timestamp("2026-01-01"), pd.Timestamp("2026-01-02 12:30:15.123456789")],
        [pd.Timestamp("2026-01-01", tz="UTC"), pd.NaT],
    ],
    ids=["midnight", "nanoseconds", "timezone"],
)
def test_csv_fidelity_preserves_timestamp_text(tmp_path: Path, values: list) -> None:
    source = pd.DataFrame({"timestamp": values})
    original = source.copy(deep=True)

    export_results(source, directory=tmp_path, filename="timestamps", extension="csv")

    restored = pd.read_csv(
        tmp_path / "timestamps.csv", dtype=str, keep_default_na=False
    )
    for text, value in zip(restored["timestamp"], values, strict=True):
        if pd.isna(value):
            assert text == ""
        else:
            assert " " in text
            assert pd.Timestamp(text) == value
    pd.testing.assert_frame_equal(source, original)


@pytest.mark.parametrize("extra_columns", [0, 7])
def test_csv_fidelity_preserves_timestamps_across_chunks(
    tmp_path: Path, extra_columns: int
) -> None:
    chunk_rows = 100_000 // (extra_columns + 1)
    values = [pd.Timestamp("2002-10-10")] * chunk_rows + [
        pd.Timestamp("2001-06-25"),
        pd.Timestamp("1993-04-28 14:40:00"),
        pd.Timestamp("1993-04-28 14:40:00.123456789"),
        pd.NaT,
    ]
    source = pd.DataFrame({"APPOINTMENT_DATETIME": values})
    for position in range(extra_columns):
        source[f"other_{position}"] = position
    original = source.copy(deep=True)

    export_results(source, directory=tmp_path, filename="appointments", extension="csv")

    restored = pd.read_csv(
        tmp_path / "appointments.csv", dtype=str, keep_default_na=False
    )
    assert restored["APPOINTMENT_DATETIME"].tolist() == [
        "2002-10-10 00:00:00"
    ] * chunk_rows + [
        "2001-06-25 00:00:00",
        "1993-04-28 14:40:00",
        "1993-04-28 14:40:00.123456789",
        "",
    ]
    pd.testing.assert_frame_equal(source, original)


def test_csv_fidelity_preserves_dates_and_nullable_integers(tmp_path: Path) -> None:
    source = pd.DataFrame(
        {
            "date": [date(2026, 1, 1), None],
            "integer": pd.Series([2**53 + 1, None], dtype="Int64"),
        }
    )

    export_results(source, directory=tmp_path, filename="values", extension="csv")

    restored = pd.read_csv(tmp_path / "values.csv", dtype=str, keep_default_na=False)
    assert restored.to_dict("list") == {
        "date": ["2026-01-01", ""],
        "integer": [str(2**53 + 1), ""],
    }


def test_xlsx_fidelity_keeps_string_data_literal(tmp_path: Path) -> None:
    source = pd.DataFrame({"text": ["=1+1", "https://example.com", "00123", "NA"]})

    export_results(source, directory=tmp_path, filename="literals", extension="xlsx")

    with (tmp_path / "literals.xlsx").open("rb") as stream:
        workbook = load_workbook(BytesIO(stream.read()))
    try:
        cells = [row[0] for row in workbook.active.iter_rows(min_row=2)]
        assert [cell.value for cell in cells] == source["text"].tolist()
        assert all(cell.data_type == "s" for cell in cells)
        assert all(cell.hyperlink is None for cell in cells)
    finally:
        workbook.close()

    restored = extract_spreadsheets(tmp_path / "literals.xlsx", dtype={"text": str})[
        0
    ].df
    pd.testing.assert_frame_equal(restored, source)


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_import_fidelity_honors_explicit_types(tmp_path: Path, extension: str) -> None:
    source = pd.DataFrame(
        {
            "float": [1.0, None, 2.0],
            "identifier": ["00123", "NA", "00456"],
            "timestamp": pd.to_datetime(["2026-01-01", None, "2026-01-02"]),
        }
    )
    export_results(source, directory=tmp_path, filename="typed", extension=extension)

    restored = extract_spreadsheets(
        tmp_path / f"typed.{extension}",
        dtype={"float": "float64", "identifier": str},
        parse_dates=["timestamp"],
    )[0].df

    pd.testing.assert_frame_equal(restored, source)


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_import_fidelity_honors_null_options(tmp_path: Path, extension: str) -> None:
    source = pd.DataFrame({"text": ["NA", "", "NULL", "nan", "other"]})
    export_results(source, directory=tmp_path, filename="nulls", extension=extension)

    restored = extract_spreadsheets(
        tmp_path / f"nulls.{extension}",
        dtype=str,
        keep_default_na=False,
        na_values=[],
    )[0].df
    pd.testing.assert_frame_equal(restored, source)

    defaults = extract_spreadsheets(tmp_path / f"nulls.{extension}", dtype=str)[0].df
    assert defaults["text"].isna().tolist() == [False, True, False, False, False]


@pytest.mark.parametrize("integer", [2**53 + 1, 2**63 - 1, 2**63, 2**64 - 1])
def test_csv_import_fidelity_preserves_large_nullable_integers(
    tmp_path: Path, integer: int
) -> None:
    source = pd.DataFrame(
        {
            "integer": pd.Series(
                [integer, None], dtype="Int64" if integer < 2**63 else "UInt64"
            ),
            "row": [0, 1],
        }
    )
    export_results(source, directory=tmp_path, filename="integers", extension="csv")

    options = {"dtype": {"integer": "UInt64"}} if integer >= 2**63 else {}
    restored = extract_spreadsheets(tmp_path / "integers.csv", **options)[0].df

    assert int(restored["integer"].iloc[0]) == integer
    assert pd.isna(restored["integer"].iloc[1])
    assert restored["integer"].dtype == source["integer"].dtype


@pytest.mark.parametrize("clean_df", [True, False])
def test_parquet_import_fidelity_preserves_schema(
    tmp_path: Path, clean_df: bool
) -> None:
    source = pd.DataFrame(
        {
            "integer": pd.Series([2**53 + 1, None], dtype="Int64"),
            "unsigned": pd.Series([2**64 - 1, None], dtype="UInt64"),
            "float": pd.Series([1.0, None], dtype="Float64"),
            "boolean": pd.Series([True, None], dtype="boolean"),
            "string": pd.Series(["00123", None], dtype="string"),
            "empty": pd.Series(["", "NA"], dtype="string"),
            "all_null": pd.Series([None, None], dtype="Int64"),
            "timestamp": pd.Series([pd.Timestamp("2026-01-01"), pd.NaT]),
            "precise_timestamp": pd.Series(
                [pd.Timestamp("2026-01-01 01:02:03.123456789", tz="UTC"), pd.NaT]
            ),
        }
    )
    export_results(source, directory=tmp_path, filename="schema", extension="parquet")

    result = extract_spreadsheets(tmp_path / "schema.parquet", clean_df=clean_df)[0]

    pd.testing.assert_frame_equal(result.df, source)
    assert result.df_name == "schema"
    assert result.filename == "schema.parquet"
    assert Path(result.filepath) == tmp_path / "schema.parquet"


@pytest.mark.parametrize("extension", ["csv", "xlsx", "parquet"])
def test_export_fidelity_omits_indexes(tmp_path: Path, extension: str) -> None:
    source = pd.DataFrame(
        {"value": [1, 2]}, index=pd.Index([10, 20], name="identifier")
    )
    original = source.copy(deep=True)
    export_results(source, directory=tmp_path, filename="indexed", extension=extension)

    restored = extract_spreadsheets(tmp_path / f"indexed.{extension}")[0].df

    assert restored.columns.tolist() == ["value"]
    assert restored.index.tolist() == [0, 1]
    pd.testing.assert_frame_equal(source, original)


def test_import_fidelity_supports_mixed_format_directories(tmp_path: Path) -> None:
    source = pd.DataFrame({"value": [1, 2]})
    for extension in ["csv", "xlsx", "parquet"]:
        export_results(
            source, directory=tmp_path, filename=extension, extension=extension
        )

    results = extract_spreadsheets(tmp_path)

    assert {result.df_name for result in results} == {"csv", "xlsx", "parquet"}
    assert all(result.df["value"].tolist() == [1, 2] for result in results)


def test_csv_import_fidelity_respects_cleaning_opt_out(tmp_path: Path) -> None:
    source = pd.DataFrame({"value": pd.Series([1, None], dtype="Int64"), "row": [0, 1]})
    export_results(source, directory=tmp_path, filename="legacy", extension="csv")

    result = extract_spreadsheets(tmp_path / "legacy.csv", clean_df=False)[0]

    assert str(result.df["value"].dtype) == "float64"


def test_csv_import_fidelity_respects_caller_backend(tmp_path: Path) -> None:
    source = pd.DataFrame({"value": [1.0, None], "row": [0, 1]})
    export_results(source, directory=tmp_path, filename="arrow", extension="csv")

    result = extract_spreadsheets(tmp_path / "arrow.csv", dtype_backend="pyarrow")[0]

    assert str(result.df["value"].dtype) == "double[pyarrow]"
    assert result.df["value"].iloc[0] == 1.0
    assert pd.isna(result.df["value"].iloc[1])


def test_csv_import_fidelity_respects_decimal_converter(tmp_path: Path) -> None:
    source = pd.DataFrame(
        {"value": [Decimal("1234567890.123456789"), None], "row": [0, 1]}
    )
    export_results(source, directory=tmp_path, filename="decimal", extension="csv")

    result = extract_spreadsheets(
        tmp_path / "decimal.csv",
        converters={"value": lambda value: Decimal(value) if value else None},
    )[0]

    assert result.df["value"].iloc[0] == source["value"].iloc[0]
    assert pd.isna(result.df["value"].iloc[1])


def test_xlsx_fidelity_distinguishes_date_and_midnight_datetime(tmp_path: Path) -> None:
    source = pd.DataFrame(
        {"date": [date(2026, 1, 1)], "timestamp": [pd.Timestamp("2026-01-01")]}
    )
    export_results(source, directory=tmp_path, filename="dates", extension="xlsx")

    workbook = load_workbook(tmp_path / "dates.xlsx")
    try:
        date_cell, timestamp_cell = list(workbook.active.iter_rows(min_row=2))[0]
        assert date_cell.number_format == "DD/MM/YYYY"
        assert timestamp_cell.number_format == "DD/MM/YYYY HH:MM:SS"
        assert timestamp_cell.value == datetime(2026, 1, 1)
    finally:
        workbook.close()


def test_xlsx_fidelity_rejects_timezone_loss(tmp_path: Path) -> None:
    source = pd.DataFrame({"timestamp": [pd.Timestamp("2026-01-01", tz="UTC")]})

    with pytest.raises(
        ValueError, match="Excel does not support datetimes with timezones"
    ):
        export_results(
            source, directory=tmp_path, filename="timezone", extension="xlsx"
        )

    assert not (tmp_path / "timezone.xlsx").exists()


def test_import_fidelity_keeps_string_analysis_compatible(tmp_path: Path) -> None:
    source = pd.DataFrame(
        {"label": ["alpha", "beta"], "date": ["2026-01-01", "2026-01-02"]}
    )
    export_results(source, directory=tmp_path, filename="analysis", extension="csv")
    restored = extract_spreadsheets(tmp_path / "analysis.csv")[0].df

    stats = examine(restored, show_results=False).stats

    assert stats.loc[stats["column_name"] == "label", "max_length"].iloc[0] == 5
    date_converter_df(restored, columns="date", readouts=False)
    assert str(restored["date"].dtype) == "datetime64[ns]"
    assert restored["date"].iloc[0] == pd.Timestamp("2026-01-01")


def test_csv_fidelity_handles_nonexistent_local_midnight(tmp_path: Path) -> None:
    timestamp = pd.Timestamp("2018-11-04 01:00:00", tz="America/Sao_Paulo")
    source = pd.DataFrame({"timestamp": [timestamp]})
    export_results(source, directory=tmp_path, filename="dst", extension="csv")

    restored = pd.read_csv(tmp_path / "dst.csv", dtype=str)

    assert pd.Timestamp(restored["timestamp"].iloc[0]) == timestamp


@pytest.mark.parametrize(
    "dtype", [str, defaultdict(lambda: str)], ids=["scalar", "defaultdict"]
)
def test_csv_import_fidelity_parses_dates_with_default_string_dtype(
    tmp_path: Path, dtype: object
) -> None:
    source = pd.DataFrame(
        {
            "identifier": ["00123", "00456"],
            "timestamp": pd.to_datetime(["2026-01-01", None]),
        }
    )
    export_results(source, directory=tmp_path, filename="strings", extension="csv")

    restored = extract_spreadsheets(
        tmp_path / "strings.csv", dtype=dtype, parse_dates=["timestamp"]
    )[0].df

    pd.testing.assert_frame_equal(restored, source)
